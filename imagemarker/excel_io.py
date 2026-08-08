"""Excel IO for ImageMarker.

Reading uses ``openpyxl`` in ``read_only`` mode and materialises only the meta
columns (``Name`` .. ``Status``) - the ~201 numeric sweep columns are never
loaded, displayed or modified.  Every parsed row remembers its 1-based
worksheet row index and the 1-based ``Status`` column index, which is what makes
the targeted, Status-only write back possible.

No tkinter import here: the module is importable and testable head-lessly.
"""

from __future__ import annotations

import os
import re
import shutil
from dataclasses import dataclass, field
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import openpyxl

from .data_model import METRIC_COLUMNS, normalize_int, normalize_status, to_float

#: Meta columns of the source workbook, in order.
META_COLUMNS: Tuple[str, ...] = (
    "Name",
    "Row",
    "Node",
    "ON",
    "OFF",
    "ON/OFF",
    "gm",
    "Vth",
    "Carrier Mobility",
    "Status",
)

#: Columns that must exist for the file to be usable.
REQUIRED_COLUMNS: Tuple[str, ...] = ("Name", "Row", "Node", "Status")

#: ``sam1`` / ``SAM 1`` / ``sam_1`` -> sample number
SAMPLE_NUMBER_PATTERN = re.compile(r"sam\s*_?(\d+)", re.IGNORECASE)

BACKUP_SUFFIX = ".backup.xlsx"


# --------------------------------------------------------------------------- #
# Errors
# --------------------------------------------------------------------------- #

class ExcelError(Exception):
    """Base class for every recoverable Excel problem."""


class ExcelFileLockedError(ExcelError):
    """The workbook is open in Excel (or otherwise not writable)."""

    def __init__(self, path: str, message: Optional[str] = None) -> None:
        self.path = path
        super().__init__(
            message
            or (
                f"Cannot write to '{os.path.basename(path)}' because the file is "
                f"locked.\nClose the file in Excel and try again."
            )
        )


class ExcelFormatError(ExcelError):
    """The workbook does not have the expected header layout."""


# --------------------------------------------------------------------------- #
# Rows
# --------------------------------------------------------------------------- #

@dataclass
class ExcelRow:
    """One data row of the source worksheet (meta columns only)."""

    name: str
    row: int
    node: int
    status: Optional[str] = None
    metrics: Dict[str, Optional[float]] = field(default_factory=dict)
    #: 1-based worksheet row index - the write-back target.
    sheet_row: int = 0
    #: 1-based worksheet column index of ``Status``.
    status_col: int = 0

    @property
    def key(self) -> Tuple[int, int]:
        return (self.row, self.node)


# --------------------------------------------------------------------------- #
# Sample-number matching helpers
# --------------------------------------------------------------------------- #

def extract_sample_number(text: Optional[str]) -> Optional[int]:
    """Extract the sample number from a folder/image/Excel name.

    ``'260619 p3meet ac 7kg 100mm, SAM 1'`` -> ``1``
    ``'20260619-P3MEEMT(1-3)_7kgf_100_sam1'`` -> ``1``
    """
    if not text:
        return None
    match = SAMPLE_NUMBER_PATTERN.search(str(text))
    if not match:
        return None
    try:
        return int(match.group(1))
    except ValueError:  # pragma: no cover - regex guarantees digits
        return None


def suggest_excel_name(image_prefix: Optional[str], names: Sequence[str]) -> Optional[str]:
    """Best-guess Excel ``Name`` for a loaded image folder.

    The image-folder prefix never equals the Excel ``Name``, so matching is done
    through the sample number.  A suggestion is only returned when exactly one
    Excel name carries the same number.
    """
    number = extract_sample_number(image_prefix)
    if number is None:
        return None
    candidates = [name for name in names if extract_sample_number(name) == number]
    if len(candidates) == 1:
        return candidates[0]
    return None


def suggest_name_mapping(
    image_prefixes: Sequence[str], names: Sequence[str]
) -> Dict[str, Optional[str]]:
    """Suggest one Excel ``Name`` per image name prefix.

    A recursive folder load can hold several samples at once, so every distinct
    image name prefix gets its own suggestion (``None`` when the sample number
    is missing or ambiguous).  The returned dict keeps the order of
    ``image_prefixes``.
    """
    return {
        prefix: suggest_excel_name(prefix, names)
        for prefix in dict.fromkeys(image_prefixes)
    }


def backup_path_for(path: str) -> str:
    """``.../data.xlsx`` -> ``.../data.backup.xlsx``."""
    stem, _ = os.path.splitext(path)
    return stem + BACKUP_SUFFIX


# --------------------------------------------------------------------------- #
# Source
# --------------------------------------------------------------------------- #

class ExcelSource:
    """A loaded ``.xlsx`` data file with targeted Status write-back."""

    def __init__(
        self,
        path: str,
        sheet_name: str,
        rows: Sequence[ExcelRow],
        status_col: int,
        columns: Optional[Dict[str, int]] = None,
    ) -> None:
        self.path = path
        self.sheet_name = sheet_name
        self.rows: List[ExcelRow] = list(rows)
        self.status_col = status_col
        self.columns: Dict[str, int] = dict(columns or {})

    # -- loading ---------------------------------------------------------- #

    @classmethod
    def load(cls, path: str, sheet_name: Optional[str] = None) -> "ExcelSource":
        """Read ``path`` (read-only, meta columns only)."""
        try:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
        except PermissionError as exc:
            raise ExcelFileLockedError(
                path,
                f"Cannot open '{os.path.basename(path)}': the file is locked by "
                f"another program.",
            ) from exc
        except ExcelError:
            raise
        except Exception as exc:  # openpyxl raises a zoo of exception types
            raise ExcelError(f"Failed to open '{os.path.basename(path)}': {exc}") from exc

        try:
            if sheet_name is None:
                if not workbook.sheetnames:
                    raise ExcelFormatError("The workbook contains no sheets.")
                sheet_name = workbook.sheetnames[0]
            elif sheet_name not in workbook.sheetnames:
                raise ExcelFormatError(f"Sheet '{sheet_name}' not found.")
            worksheet = workbook[sheet_name]

            columns = cls._read_header(worksheet)
            missing = [name for name in REQUIRED_COLUMNS if name not in columns]
            if missing:
                raise ExcelFormatError(
                    "Missing required column(s): " + ", ".join(missing)
                )

            rows = cls._read_rows(worksheet, columns)
        finally:
            workbook.close()

        return cls(
            path=path,
            sheet_name=sheet_name,
            rows=rows,
            status_col=columns["Status"],
            columns=columns,
        )

    @staticmethod
    def _read_header(worksheet: Any) -> Dict[str, int]:
        """Map meta column name -> 1-based column index (first match wins)."""
        try:
            header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
        except StopIteration:
            raise ExcelFormatError("The worksheet is empty.") from None

        wanted = {name.casefold(): name for name in META_COLUMNS}
        columns: Dict[str, int] = {}
        for index, value in enumerate(header, start=1):
            if value is None:
                continue
            key = str(value).strip().casefold()
            canonical = wanted.get(key)
            if canonical is not None and canonical not in columns:
                columns[canonical] = index
        return columns

    @staticmethod
    def _read_rows(worksheet: Any, columns: Dict[str, int]) -> List[ExcelRow]:
        name_col = columns["Name"]
        row_col = columns["Row"]
        node_col = columns["Node"]
        status_col = columns["Status"]
        metric_cols = {
            metric: columns[metric] for metric in METRIC_COLUMNS if metric in columns
        }
        max_col = max(columns.values())

        rows: List[ExcelRow] = []
        iterator = worksheet.iter_rows(min_row=2, max_col=max_col, values_only=True)
        for sheet_row, values in enumerate(iterator, start=2):
            if not values or all(value is None for value in values):
                continue

            def cell(index: int) -> Any:
                return values[index - 1] if index <= len(values) else None

            row_value = normalize_int(cell(row_col))
            node_value = normalize_int(cell(node_col))
            if row_value is None or node_value is None:
                continue

            raw_name = cell(name_col)
            name = "" if raw_name is None else str(raw_name).strip()

            rows.append(
                ExcelRow(
                    name=name,
                    row=row_value,
                    node=node_value,
                    status=normalize_status(cell(status_col)),
                    metrics={
                        metric: to_float(cell(index))
                        for metric, index in metric_cols.items()
                    },
                    sheet_row=sheet_row,
                    status_col=status_col,
                )
            )
        return rows

    # -- queries ---------------------------------------------------------- #

    def names(self) -> List[Tuple[str, int]]:
        """Distinct ``Name`` values with row counts, in first-seen order."""
        counts: Dict[str, int] = {}
        for row in self.rows:
            counts[row.name] = counts.get(row.name, 0) + 1
        return list(counts.items())

    def rows_for_name(self, name: str) -> List[ExcelRow]:
        return [row for row in self.rows if row.name == name]

    def index_for_name(self, name: str) -> Dict[Tuple[int, int], ExcelRow]:
        """``(row, node)`` -> row, restricted to one ``Name`` (first wins)."""
        index: Dict[Tuple[int, int], ExcelRow] = {}
        for row in self.rows_for_name(name):
            index.setdefault(row.key, row)
        return index

    def unique_statuses(self, name: Optional[str] = None) -> List[Optional[str]]:
        """Distinct status values present (dynamic, never hardcoded)."""
        rows = self.rows if name is None else self.rows_for_name(name)
        seen: List[Optional[str]] = []
        for row in rows:
            if row.status not in seen:
                seen.append(row.status)
        return seen

    def name_list(self) -> List[str]:
        """Distinct ``Name`` values, in first-seen order."""
        return [name for name, _ in self.names()]

    def suggest_name(self, image_prefix: Optional[str]) -> Optional[str]:
        return suggest_excel_name(image_prefix, self.name_list())

    def suggest_mapping(
        self, image_prefixes: Sequence[str]
    ) -> Dict[str, Optional[str]]:
        """Auto-suggest an Excel ``Name`` for every loaded image name prefix."""
        return suggest_name_mapping(image_prefixes, self.name_list())

    # -- write back -------------------------------------------------------- #

    @property
    def backup_path(self) -> str:
        return backup_path_for(self.path)

    def ensure_backup(self) -> bool:
        """Create the one-time ``.backup.xlsx`` copy; ``True`` if created now."""
        destination = self.backup_path
        if os.path.exists(destination):
            return False
        try:
            shutil.copy2(self.path, destination)
        except PermissionError as exc:
            raise ExcelFileLockedError(
                self.path,
                f"Cannot create the backup '{os.path.basename(destination)}': "
                f"permission denied.\nClose the file in Excel and try again.",
            ) from exc
        except OSError as exc:
            raise ExcelError(f"Failed to create backup: {exc}") from exc
        return True

    def write_statuses(
        self, updates: Sequence[Tuple[int, int, Optional[str]]]
    ) -> Tuple[int, bool]:
        """Write ONLY the given Status cells back into the original workbook.

        ``updates`` is a sequence of ``(sheet_row, status_col, value)``.  The
        workbook is re-opened in normal mode so all formatting, formulas and the
        ~201 sweep columns survive untouched; only the listed cells are set.

        Returns ``(written_cells, backup_created)``.
        Raises :class:`ExcelFileLockedError` when the file is open in Excel.
        """
        if not updates:
            return 0, False

        backup_created = self.ensure_backup()

        try:
            workbook = openpyxl.load_workbook(self.path)
        except PermissionError as exc:
            raise ExcelFileLockedError(self.path) from exc
        except Exception as exc:
            raise ExcelError(
                f"Failed to open '{os.path.basename(self.path)}': {exc}"
            ) from exc

        try:
            if self.sheet_name not in workbook.sheetnames:
                raise ExcelFormatError(
                    f"Sheet '{self.sheet_name}' no longer exists in the workbook."
                )
            worksheet = workbook[self.sheet_name]
            for sheet_row, status_col, value in updates:
                worksheet.cell(row=sheet_row, column=status_col).value = (
                    normalize_status(value)
                )
            try:
                workbook.save(self.path)
            except PermissionError as exc:
                raise ExcelFileLockedError(self.path) from exc
            except OSError as exc:
                raise ExcelError(f"Failed to save the workbook: {exc}") from exc
        finally:
            workbook.close()

        # Keep the in-memory rows consistent with the file.
        by_position = {(row.sheet_row, row.status_col): row for row in self.rows}
        for sheet_row, status_col, value in updates:
            row = by_position.get((sheet_row, status_col))
            if row is not None:
                row.status = normalize_status(value)

        return len(updates), backup_created

    def write_records(self, records: Iterable[Any]) -> Tuple[int, bool, int]:
        """Write back the Status of every record that knows its worksheet cell.

        Records are duck typed (``excel_row``, ``excel_status_col``, ``status``)
        so that :mod:`imagemarker.data_model` stays decoupled from this module.

        Returns ``(written_cells, backup_created, skipped_records)``.
        """
        updates: List[Tuple[int, int, Optional[str]]] = []
        skipped = 0
        for record in records:
            sheet_row = getattr(record, "excel_row", None)
            status_col = getattr(record, "excel_status_col", None)
            if sheet_row is None or status_col is None:
                skipped += 1
                continue
            updates.append((sheet_row, status_col, getattr(record, "status", None)))
        written, backup_created = self.write_statuses(updates)
        return written, backup_created, skipped
